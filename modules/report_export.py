"""
Word 经营分析报告导出模块
========================
使用 python-docx 生成格式化的经营分析报告，支持包含指标、异常、AI报告内容。
"""

import os
from datetime import datetime


def _get_report_data(metrics, anomalies, field_mapping, ai_report=None):
    """
    整理报告数据，供导出使用。

    参数:
        metrics: dict — 核心指标
        anomalies: list — 异常列表
        field_mapping: dict — 字段映射
        ai_report: str, optional — AI 生成的经营分析报告

    返回:
        dict — 整理好的报告数据
    """
    keys_present = metrics.get("_keys_present", [])

    # 整理指标
    metric_items = []
    standard_metrics = [
        ("总销售额 (GMV)", "元"),
        ("总订单量", "单"),
        ("平均客单价", "元/单"),
        ("总访客数 (UV)", "人"),
        ("总浏览量 (PV)", "次"),
        ("总退款金额", "元"),
        ("退款率", "%"),
        ("总推广花费", "元"),
        ("ROI (投产比)", ""),
        ("转化率", "%"),
    ]
    for name, unit in standard_metrics:
        if name in metrics:
            val = metrics[name]
            metric_items.append({"name": name, "value": val, "unit": unit})

    # 整理异常
    anomaly_items = []
    for a in anomalies:
        if a["type"] == "暂无异常":
            continue
        anomaly_items.append({
            "type": a["type"],
            "severity": a["severity"],
            "description": a["description"],
            "possible_causes": a.get("possible_causes", []),
            "suggest_metrics": a.get("suggest_metrics", ""),
        })

    # 整理字段信息
    field_info = []
    for std_field, original_col in field_mapping.items():
        if original_col:
            field_info.append(f"{std_field} → {original_col}")

    return {
        "metric_items": metric_items,
        "anomaly_items": anomaly_items,
        "field_info": field_info,
        "ai_report": ai_report,
        "keys_present": keys_present,
    }


def generate_basic_report(metrics, anomalies, field_mapping):
    """
    生成基础经营分析报告文本（无 AI 时使用）。

    参数:
        metrics: dict — 核心指标
        anomalies: list — 异常列表
        field_mapping: dict — 字段映射

    返回:
        str — 报告文本
    """
    data = _get_report_data(metrics, anomalies, field_mapping)
    lines = []

    # 整体概况
    lines.append("【整体概况】")
    total_gmv = metrics.get("总销售额 (GMV)", 0)
    total_orders = metrics.get("总订单量", 0)
    lines.append(f"本期总销售额为 {total_gmv:,.0f} 元，总订单量为 {total_orders:,.0f} 单。")

    avg_price = metrics.get("平均客单价", 0)
    if avg_price:
        lines.append(f"平均客单价为 {avg_price:.2f} 元/单。")

    refund_rate = metrics.get("退款率", 0)
    if refund_rate:
        lines.append(f"整体退款率为 {refund_rate*100:.2f}%。")

    roi = metrics.get("ROI (投产比)", 0)
    if roi:
        lines.append(f"整体投产比为 {roi:.2f}。")

    conversion = metrics.get("转化率", 0)
    if conversion:
        lines.append(f"整体转化率为 {conversion*100:.2f}%。")

    # 核心指标
    lines.append("")
    lines.append("【核心指标】")
    for item in data["metric_items"]:
        name = item["name"]
        val = item["value"]
        unit = item["unit"]
        if unit == "%":
            # 退款率和转化率需要特殊格式化
            if "退款率" in name:
                lines.append(f"  {name}：{val*100:.2f}%")
            elif "转化率" in name:
                lines.append(f"  {name}：{val*100:.2f}%")
            else:
                lines.append(f"  {name}：{val:,.2f}{unit}")
        elif unit == "" and "ROI" in name:
            lines.append(f"  {name}：{val:.2f}")
        else:
            lines.append(f"  {name}：{val:,.0f} {unit}")

    # 异常识别
    if data["anomaly_items"]:
        lines.append("")
        lines.append("【异常识别】")
        for a in data["anomaly_items"]:
            sev_map = {"高": "🔴 高风险", "中": "🟡 中风险", "低": "🟢 低风险"}
            sev_label = sev_map.get(a["severity"], a["severity"])
            lines.append(f"  {sev_label}：{a['type']}")
            lines.append(f"    描述：{a['description']}")
            if a["possible_causes"]:
                for cause in a["possible_causes"][:2]:  # 最多2条
                    lines.append(f"    可能原因：{cause}")
            lines.append("")
    else:
        lines.append("")
        lines.append("【异常识别】\n  未检测到明显的经营异常情况。")

    # 经营建议
    lines.append("【经营建议】")
    if metrics.get("退款率", 0) and metrics.get("退款率", 0) > 0.1:
        lines.append("  1. 退款率偏高（超过 10%），建议排查退款原因，优化商品质量和售后流程。")
    if metrics.get("ROI (投产比)", 0) and 0 < metrics["ROI (投产比)"] < 2:
        lines.append("  2. ROI 偏低，建议优化推广投放策略，提高投产比。")
    if metrics.get("总订单量", 0) and metrics.get("总访客数 (UV)", 0):
        conv = metrics.get("转化率", 0)
        if conv and conv < 0.05:
            lines.append("  3. 转化率偏低，建议优化商品详情页、定价策略和促销活动。")
    if not any("建议" in l for l in lines[-5:]):  # 没有具体建议时
        lines.append("  当前经营数据整体表现平稳，建议关注各指标变化趋势，保持健康增长。")
    lines.append("  4. 建议持续监控销售额、订单量和退款率等核心指标的变化趋势。")

    return "\n".join(lines)


def export_report_to_docx(metrics, anomalies, field_mapping, ai_report=None):
    """
    导出经营分析报告为 Word 文档。

    参数:
        metrics: dict — 核心指标
        anomalies: list — 异常列表
        field_mapping: dict — 字段映射
        ai_report: str, optional — AI 生成的经营分析报告文本

    返回:
        success: bool
        result: bytes or str — 文档字节流或错误信息
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
    except ImportError:
        return False, "请先安装 python-docx：pip install python-docx"

    try:
        doc = Document()

        # ---- 全局样式设置 ----
        style = doc.styles["Normal"]
        font = style.font
        font.name = "微软雅黑"
        font.size = Pt(10.5)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

        # ---- 标题 ----
        title = doc.add_heading("企业经营分析报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 副标题：生成日期
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

        doc.add_paragraph()  # 空行

        # ---- 一、整体概况 ----
        doc.add_heading("一、整体经营概况", level=1)

        total_gmv = metrics.get("总销售额 (GMV)", 0)
        total_orders = metrics.get("总订单量", 0)
        p = doc.add_paragraph()
        p.add_run(f"本期总销售额为 {total_gmv:,.0f} 元，总订单量为 {total_orders:,.0f} 单。")

        avg_price = metrics.get("平均客单价", 0)
        if avg_price:
            p = doc.add_paragraph()
            p.add_run(f"平均客单价为 {avg_price:.2f} 元/单。")

        refund_rate = metrics.get("退款率", 0)
        if refund_rate:
            p = doc.add_paragraph()
            p.add_run(f"整体退款率为 {refund_rate*100:.2f}%。")

        roi = metrics.get("ROI (投产比)", 0)
        if roi:
            p = doc.add_paragraph()
            p.add_run(f"整体投产比为 {roi:.2f}。")

        conversion = metrics.get("转化率", 0)
        if conversion:
            p = doc.add_paragraph()
            p.add_run(f"整体转化率为 {conversion*100:.2f}%。")

        # ---- 二、核心指标 ----
        doc.add_heading("二、核心经营指标", level=1)

        data = _get_report_data(metrics, anomalies, field_mapping)
        if data["metric_items"]:
            # 指标表格（2 列：指标名称、数值）
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Shading Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            hdr = table.rows[0].cells
            hdr[0].text = "指标名称"
            hdr[1].text = "数值"

            keys_present = metrics.get("_keys_present", [])
            for item in data["metric_items"]:
                row_cells = table.add_row().cells
                row_cells[0].text = item["name"]
                val = item["value"]
                unit = item["unit"]
                name = item["name"]

                if unit == "%":
                    if "退款率" in name or "转化率" in name:
                        row_cells[1].text = f"{val*100:.2f}%"
                    else:
                        row_cells[1].text = f"{val:,.2f}{unit}"
                elif unit == "" and "ROI" in name:
                    row_cells[1].text = f"{val:.2f}"
                else:
                    row_cells[1].text = f"{val:,.0f} {unit}".strip()

            # 设置表格列宽
            for row in table.rows:
                row.cells[0].width = Cm(6)
                row.cells[1].width = Cm(6)
        else:
            doc.add_paragraph("无可用指标数据。")

        # ---- 三、异常识别 ----
        doc.add_heading("三、经营异常识别", level=1)

        if data["anomaly_items"]:
            sev_map = {"高": "🔴 高风险", "中": "🟡 中风险", "低": "🟢 低风险"}
            for a in data["anomaly_items"]:
                sev_label = sev_map.get(a["severity"], a["severity"])
                p = doc.add_paragraph()
                run = p.add_run(f"{sev_label}：{a['type']}")
                run.bold = True

                p = doc.add_paragraph()
                p.add_run(f"描述：{a['description']}")

                if a["possible_causes"]:
                    p = doc.add_paragraph()
                    p.add_run("可能原因：")
                    for cause in a["possible_causes"]:
                        doc.add_paragraph(cause, style="List Bullet")

                if a["suggest_metrics"]:
                    p = doc.add_paragraph()
                    p.add_run(f"建议关注指标：{a['suggest_metrics']}")

                doc.add_paragraph()  # 行间距
        else:
            doc.add_paragraph("未检测到明显的经营异常情况。")

        # ---- 四、经营建议 ----
        doc.add_heading("四、经营优化建议", level=1)

        suggestions = []
        if metrics.get("退款率", 0) and metrics["退款率"] > 0.1:
            suggestions.append(
                "退款率偏高（超过 10%），建议排查退款原因：检查商品质量、"
                "优化详情页描述、改进物流服务和售后流程。"
            )
        if metrics.get("ROI (投产比)", 0) and 0 < metrics["ROI (投产比)"] < 2:
            suggestions.append(
                "ROI 偏低，建议优化推广投放策略：调整出价、优化人群定向、"
                "测试不同素材和落地页效果。"
            )
        if metrics.get("转化率", 0) and metrics["转化率"] < 0.05:
            suggestions.append(
                "转化率偏低，建议分析转化漏斗各环节流失情况，"
                "优化商品详情页、定价策略和促销活动。"
            )

        if not suggestions:
            suggestions.append(
                "当前经营数据整体表现平稳，建议持续关注各指标变化趋势，"
                "保持健康增长态势。"
            )
        suggestions.append(
            "建议持续监控销售额、订单量、退款率和 ROI 等核心指标的变化趋势，"
            "定期进行经营分析诊断。"
        )

        for s in suggestions:
            doc.add_paragraph(s, style="List Number")

        # ---- 五、AI 分析报告（如有） ----
        if ai_report:
            doc.add_page_break()
            doc.add_heading("五、AI 智能分析补充", level=1)
            p = doc.add_paragraph()
            p.add_run("以下内容由 AI 基于数据分析生成，供参考：")
            doc.add_paragraph()

            # 将 AI 报告的 Markdown 分行添加到文档
            for line in ai_report.strip().split("\n"):
                line = line.strip()
                if not line:
                    continue
                # 判断是否为大标题（以 ## 开头）
                if line.startswith("## "):
                    doc.add_heading(line.replace("## ", "").strip(), level=2)
                elif line.startswith("# "):
                    doc.add_heading(line.replace("# ", "").strip(), level=1)
                elif line.startswith("**") and line.endswith("**"):
                    doc.add_heading(line.strip("*"), level=3)
                elif line.startswith("- ") or line.startswith("* "):
                    doc.add_paragraph(line[2:], style="List Bullet")
                elif any(line.startswith(f"{i}.") for i in range(1, 10)):
                    doc.add_paragraph(line, style="List Number")
                else:
                    doc.add_paragraph(line)

        # ---- 页脚信息 ----
        doc.add_paragraph()
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_p.add_run(
            "— 本报告由 AI 企业经营分析 Agent 自动生成 —"
        )
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(180, 180, 180)

        # 保存到字节流
        from io import BytesIO
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return True, buf.getvalue()

    except Exception as e:
        return False, f"报告导出失败：{str(e)}"


def generate_summary_brief(metrics, anomalies):
    """
    生成简短经营概况文本（适合展示在按钮旁或报告头部）。

    参数:
        metrics: dict — 核心指标
        anomalies: list — 异常列表

    返回:
        str — 简短概况
    """
    lines = []
    total_gmv = metrics.get("总销售额 (GMV)", 0)
    total_orders = metrics.get("总订单量", 0)
    lines.append(f"本期总销售额 {total_gmv:,.0f} 元，总订单量 {total_orders:,.0f} 单")

    refund_rate = metrics.get("退款率", 0)
    if refund_rate:
        lines.append(f"退款率 {refund_rate*100:.2f}%")

    roi = metrics.get("ROI (投产比)", 0)
    if roi:
        lines.append(f"ROI {roi:.2f}")

    # 统计异常数量（排除"暂无异常"）
    real_anomalies = [a for a in anomalies if a["type"] != "暂无异常"]
    if real_anomalies:
        high = sum(1 for a in real_anomalies if a["severity"] == "高")
        lines.append(f"发现 {len(real_anomalies)} 项异常（高风险 {high} 项）")
    else:
        lines.append("未发现明显异常")

    return " | ".join(lines)
